import io
import re
from typing import List

from pypdf import PdfReader
from docx import Document as DocxDocument

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


# ============================================================
# DOCUMIND AI - RAG ENGINE
# ============================================================
#
# Responsibilities:
#
# 1. Load PDF / DOCX / TXT
# 2. Preserve source metadata
# 3. Preserve PDF page numbers
# 4. Clean extracted text
# 5. Split documents into chunks
# 6. Preserve metadata after chunking
#
# Future RAG pipeline:
#
# Documents
#     ↓
# Page-aware extraction
#     ↓
# Chunking
#     ↓
# Embeddings → FAISS
#     +
# BM25
#     ↓
# Fusion / RRF
#     ↓
# Reranking
#     ↓
# LLM
#     ↓
# Answer + Page Citations
#
# ============================================================


# ============================================================
# CONFIGURATION
# ============================================================

CHUNK_SIZE = 800

CHUNK_OVERLAP = 120


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(text: str) -> str:
    """
    Clean extracted document text.

    Removes excessive spaces and blank lines while
    preserving the actual document content.
    """

    if not text:
        return ""

    # Normalize carriage returns
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Replace tabs with spaces
    text = text.replace("\t", " ")

    # Remove excessive spaces
    text = re.sub(
        r"[ ]{2,}",
        " ",
        text
    )

    # Remove excessive blank lines
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# PDF LOADER
# ============================================================

def load_pdf(
    file_bytes: bytes,
    filename: str
) -> List[Document]:
    """
    Extract text from every PDF page.

    IMPORTANT:
    Page metadata is preserved so that later we can
    show citations such as:

        document.pdf — Page 247
    """

    documents = []

    pdf_stream = io.BytesIO(
        file_bytes
    )

    print()
    print("=" * 60)
    print(f"Loading PDF: {filename}")
    print("=" * 60)

    try:

        reader = PdfReader(
            pdf_stream
        )

    except Exception as e:

        raise RuntimeError(
            f"Unable to read PDF '{filename}': {e}"
        )

    total_pages = len(
        reader.pages
    )

    print(
        f"Total pages: {total_pages}"
    )

    print()

    # --------------------------------------------------------
    # PROCESS EACH PAGE
    # --------------------------------------------------------

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        try:

            text = page.extract_text()

            text = clean_text(
                text
            )

            # ------------------------------------------------
            # Ignore completely empty pages
            # ------------------------------------------------

            if text:

                document = Document(

                    page_content=text,

                    metadata={
                        "source": filename,
                        "file_name": filename,
                        "file_type": "pdf",

                        # VERY IMPORTANT
                        "page": page_number,

                        "page_number": page_number,

                        "total_pages": total_pages,
                    }
                )

                documents.append(
                    document
                )

        except Exception as e:

            print(
                f"Warning: "
                f"Could not process page "
                f"{page_number}: {e}"
            )

        # ----------------------------------------------------
        # PROGRESS
        # ----------------------------------------------------

        if (
            page_number % 10 == 0
            or page_number == total_pages
        ):

            percentage = (
                page_number
                / total_pages
                * 100
            )

            print(
                f"Processed "
                f"{page_number}/{total_pages} "
                f"pages "
                f"({percentage:.1f}%)"
            )

    print()

    print(
        f"Successfully extracted "
        f"{len(documents)} pages with text."
    )

    return documents


# ============================================================
# DOCX LOADER
# ============================================================

def load_docx(
    file_bytes: bytes,
    filename: str
) -> List[Document]:
    """
    Extract text from a DOCX document.

    DOCX does not have reliable page metadata at the
    paragraph level, so page is set to None.
    """

    documents = []

    doc_stream = io.BytesIO(
        file_bytes
    )

    try:

        doc = DocxDocument(
            doc_stream
        )

    except Exception as e:

        raise RuntimeError(
            f"Unable to read DOCX "
            f"'{filename}': {e}"
        )

    paragraphs = []

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:

            paragraphs.append(
                text
            )

    full_text = "\n\n".join(
        paragraphs
    )

    full_text = clean_text(
        full_text
    )

    if full_text:

        document = Document(

            page_content=full_text,

            metadata={
                "source": filename,
                "file_name": filename,
                "file_type": "docx",

                "page": None,

                "page_number": None,

                "total_pages": None,
            }
        )

        documents.append(
            document
        )

    return documents


# ============================================================
# TXT LOADER
# ============================================================

def load_txt(
    file_bytes: bytes,
    filename: str
) -> List[Document]:
    """
    Load a TXT document.
    """

    documents = []

    try:

        text = file_bytes.decode(
            "utf-8",
            errors="ignore"
        )

    except Exception as e:

        raise RuntimeError(
            f"Unable to read TXT "
            f"'{filename}': {e}"
        )

    text = clean_text(
        text
    )

    if text:

        document = Document(

            page_content=text,

            metadata={
                "source": filename,
                "file_name": filename,
                "file_type": "txt",

                "page": None,

                "page_number": None,

                "total_pages": None,
            }
        )

        documents.append(
            document
        )

    return documents


# ============================================================
# GENERIC DOCUMENT LOADER
# ============================================================

def load_document(
    file_bytes: bytes,
    filename: str
) -> List[Document]:
    """
    Automatically select the correct loader
    based on file extension.
    """

    extension = (
        filename
        .lower()
        .split(".")[-1]
    )

    if extension == "pdf":

        return load_pdf(
            file_bytes,
            filename
        )

    elif extension == "docx":

        return load_docx(
            file_bytes,
            filename
        )

    elif extension == "txt":

        return load_txt(
            file_bytes,
            filename
        )

    else:

        raise ValueError(
            f"Unsupported file type: "
            f".{extension}"
        )


# ============================================================
# CHUNKING
# ============================================================

def create_chunks(
    documents: List[Document]
) -> List[Document]:
    """
    Split documents into smaller chunks.

    IMPORTANT:
    LangChain preserves metadata when splitting
    Documents.

    Therefore:

        source
        page
        file_type
        page_number

    remain attached to each chunk.
    """

    print()
    print("=" * 60)
    print("CREATING DOCUMENT CHUNKS")
    print("=" * 60)

    print(
        f"Input documents/pages: "
        f"{len(documents)}"
    )

    print(
        f"Chunk size: {CHUNK_SIZE}"
    )

    print(
        f"Chunk overlap: {CHUNK_OVERLAP}"
    )

    print()

    splitter = (
        RecursiveCharacterTextSplitter(

            chunk_size=CHUNK_SIZE,

            chunk_overlap=CHUNK_OVERLAP,

            separators=[
                "\n\n",
                "\n",
                ". ",
                "? ",
                "! ",
                " ",
                "",
            ],
        )
    )

    chunks = splitter.split_documents(
        documents
    )

    # --------------------------------------------------------
    # Add unique chunk IDs
    # --------------------------------------------------------

    for index, chunk in enumerate(
        chunks
    ):

        chunk.metadata[
            "chunk_id"
        ] = index

    print(
        f"Total chunks created: "
        f"{len(chunks)}"
    )

    return chunks


# ============================================================
# PROCESS ONE FILE
# ============================================================

def process_file(file) -> List[Document]:
    """
    Process a single Streamlit UploadedFile.

    Example:

        chunks = process_file(uploaded_file)
    """

    filename = file.name

    file_bytes = file.getvalue()

    print()
    print("=" * 60)
    print(f"PROCESSING FILE: {filename}")
    print("=" * 60)

    print(
        f"File size: "
        f"{len(file_bytes) / (1024 * 1024):.2f} MB"
    )

    documents = load_document(
        file_bytes,
        filename
    )

    chunks = create_chunks(
        documents
    )

    return chunks


# ============================================================
# PROCESS MULTIPLE FILES
# ============================================================

def process_files(
    files
) -> List[Document]:
    """
    Process multiple uploaded files.

    This is the function that your Streamlit
    application will eventually call.
    """

    all_chunks = []

    print()
    print("#" * 60)
    print("DOCUMIND AI - PROCESSING DOCUMENTS")
    print("#" * 60)

    print(
        f"Files received: {len(files)}"
    )

    print()

    # --------------------------------------------------------
    # Process every file
    # --------------------------------------------------------

    for file_number, file in enumerate(
        files,
        start=1
    ):

        print(
            f"[{file_number}/{len(files)}] "
            f"{file.name}"
        )

        chunks = process_file(
            file
        )

        all_chunks.extend(
            chunks
        )

    # --------------------------------------------------------
    # Reassign global chunk IDs
    # --------------------------------------------------------

    for index, chunk in enumerate(
        all_chunks
    ):

        chunk.metadata[
            "chunk_id"
        ] = index

    print()
    print("#" * 60)
    print("PROCESSING COMPLETE")
    print("#" * 60)

    print(
        f"Total chunks: "
        f"{len(all_chunks)}"
    )

    return all_chunks


# ============================================================
# DEBUG / INSPECTION HELPER
# ============================================================

def inspect_chunks(
    chunks: List[Document],
    number: int = 5
):
    """
    Print a small sample of chunks.

    Useful while developing the RAG pipeline.
    """

    print()
    print("=" * 60)
    print("CHUNK INSPECTION")
    print("=" * 60)

    for chunk in chunks[:number]:

        print()

        print(
            "Chunk ID:",
            chunk.metadata.get(
                "chunk_id"
            )
        )

        print(
            "Source:",
            chunk.metadata.get(
                "source"
            )
        )

        print(
            "File type:",
            chunk.metadata.get(
                "file_type"
            )
        )

        print(
            "Page:",
            chunk.metadata.get(
                "page"
            )
        )

        print(
            "Characters:",
            len(
                chunk.page_content
            )
        )

        print(
            "Content:"
        )

        print(
            chunk.page_content[:300]
        )

        print(
            "-" * 60
        )


# ============================================================
# END
# ============================================================